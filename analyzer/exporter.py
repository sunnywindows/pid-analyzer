"""分析结果导出为 Excel。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _write_sheet(ws, title: str, rows: list[dict], columns: list[tuple[str, str]]):
    ws.title = title
    ws.append([c[1] for c in columns])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r.get(c[0], "") for c in columns])
    # 列宽自适应（粗略）
    for i, (key, _label) in enumerate(columns, start=1):
        max_len = max([len(str(r.get(key, ""))) for r in rows] + [len(_label)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len * 2 + 2, 10), 60)
    ws.freeze_panes = "A2"


def export_to_excel(result: dict) -> bytes:
    """把统一分析结果导出为 xlsx 字节流。"""
    wb = Workbook()

    ws = wb.active
    ws.title = "概览"
    meta = result.get("meta", {})
    overview = [
        ["文件名称", meta.get("filename", "")],
        ["分析引擎", "AI 视觉大模型" if meta.get("engine") == "ai" else "本地 OCR 规则"],
        ["模型", meta.get("model", "—")],
        ["页数", meta.get("pages", 1)],
        ["图纸尺寸", meta.get("dimensions", "—")],
        ["分析时间", meta.get("analyzed_at", "")],
        ["", ""],
        ["图纸概要", result.get("summary", "")],
        ["物料流向", result.get("flow", "")],
    ]
    for row in overview:
        ws.append(row)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90
    for r in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=1):
        r[0].font = Font(bold=True)

    _write_sheet(wb.create_sheet(), "设备清单", result.get("equipment", []),
                 [("tag", "位号"), ("type", "类型"), ("description", "说明"), ("confidence", "置信度")])
    _write_sheet(wb.create_sheet(), "仪表清单", result.get("instruments", []),
                 [("tag", "位号"), ("type", "类型"), ("description", "说明"), ("confidence", "置信度")])
    _write_sheet(wb.create_sheet(), "管线清单", result.get("pipelines", []),
                 [("tag", "管线号"), ("type", "介质/类型"), ("description", "起终点/规格"), ("confidence", "置信度")])

    if result.get("findings"):
        ws = wb.create_sheet("发现与建议")
        ws.append(["序号", "内容"])
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, f in enumerate(result["findings"], start=1):
            ws.append([i, f])
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
