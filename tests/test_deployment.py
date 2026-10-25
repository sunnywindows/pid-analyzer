import io
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

import app as application
from analyzer import ai_analyzer


class DeploymentTests(unittest.TestCase):
    def setUp(self):
        application._tasks.clear()
        self.client = TestClient(application.app)
        self.result = {
            "summary": "测试图纸概要",
            "flow": "V-101 至 P-101",
            "equipment": [{"tag": "V-101", "type": "容器", "description": "测试设备"}],
            "instruments": [{"tag": "PI-101", "type": "压力指示"}],
            "pipelines": [{"tag": "L-101", "from": "V-101", "to": "P-101"}],
            "findings": ["请复核安全阀位号。"],
        }

    def model_response(self):
        import json

        response = Mock()
        response.json.return_value = {
            "choices": [{"message": {"content": json.dumps(self.result, ensure_ascii=False)}}]
        }
        return response

    def submit(self, model="gpt-6-luna"):
        sample = Path(__file__).resolve().parents[1] / "sample" / "HP-Gas-Separator-PID.png"
        response = self.client.post(
            "/api/analyze",
            data={"engine": "ai", "api_key": "test-key-not-real", "model": model},
            files={"file": (sample.name, sample.read_bytes(), "image/png")},
        )
        self.assertEqual(response.status_code, 200)
        task_id = response.json()["task_id"]
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            status = self.client.get(f"/api/status/{task_id}").json()
            if status["status"] != "running":
                return status
            time.sleep(0.05)
        self.fail("后台分析任务没有在规定时间内完成")

    def test_home_and_static_files(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn('window.PID_MODE = "ai"', response.text)
        self.assertEqual(self.client.get("/static/app.js").status_code, 200)

    def test_gpt6_upload_task_and_export(self):
        with patch("analyzer.ai_analyzer.requests.post", return_value=self.model_response()) as post:
            status = self.submit()
        self.assertEqual(status["status"], "done", status)
        result = status["result"]
        self.assertEqual(result["equipment"][0]["tag"], "V-101")
        self.assertEqual(result["findings"], self.result["findings"])
        self.assertTrue(result["preview"])
        self.assertEqual(post.call_args.args[0], "https://api.openai.com/v1/chat/completions")
        payload = post.call_args.kwargs["json"]
        self.assertEqual(payload["model"], "gpt-6-luna")
        self.assertEqual(payload["max_completion_tokens"], 4096)
        self.assertNotIn("max_tokens", payload)
        self.assertNotIn("temperature", payload)
        self.assertTrue(payload["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))

        exported = self.client.post("/api/export", json=result)
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
        try:
            self.assertEqual(len(workbook.sheetnames), 5)
            values = [value for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
            self.assertIn("V-101", values)
            self.assertIn(self.result["findings"][0], values)
        finally:
            workbook.close()

    def test_other_provider_request_is_preserved(self):
        with patch("analyzer.ai_analyzer.requests.post", return_value=self.model_response()) as post:
            ai_analyzer.analyze(
                [Image.new("RGB", (32, 32), "white")],
                api_key="test-key-not-real",
                model="qwen-vl-max",
                base_url="https://compatible-provider.invalid/v1",
            )
        payload = post.call_args.kwargs["json"]
        self.assertEqual(payload["max_tokens"], 4096)
        self.assertEqual(payload["temperature"], 0.1)
        self.assertNotIn("max_completion_tokens", payload)

    def test_connection_failure_reaches_task_status(self):
        import requests

        with patch("analyzer.ai_analyzer.requests.post", side_effect=requests.ConnectionError("Network is unreachable")):
            status = self.submit()
        self.assertEqual(status["status"], "error")
        self.assertIn("Network is unreachable", status["error"])

    def test_invalid_uploads_are_rejected(self):
        for name, content in [("unsupported.txt", b"text"), ("empty.png", b"")]:
            response = self.client.post("/api/analyze", files={"file": (name, content)})
            self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
